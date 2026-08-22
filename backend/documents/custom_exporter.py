import os
import re
import json
import copy
import zipfile
import traceback
import threading
import time
import uuid
from contextlib import contextmanager
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor, as_completed
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Inches
from datetime import datetime, timezone

from backend.documents.archive_validation import validate_ooxml_archive
from backend.documents.docx_column_loop import (
    COLUMN_LITERAL_CONTEXT,
    expand_column_loops,
)
from backend.documents.template_security import (
    create_template_environment,
    validate_template_root_keys,
    validate_template_statements,
)
from backend.documents.docx_context_policy import (
    BASE_IMAGE_FIELDS,
    validate_docx_context_manifest,
)
from backend.shared.paths import IMAGE_DIR, PROJECT_ROOT, WORD_TEMPLATE_DIR
from backend.shared.logging_utils import append_runtime_log, log_error
from backend.shared.media_helper import (
    managed_image_path_matches_tenant,
    normalize_managed_image_path,
)


_IMAGE_THREAD_POOL = ThreadPoolExecutor(max_workers=6)
_WORD_CONFIG_LOCK = threading.RLock()
_WORD_CONFIG_LOCK_TIMEOUT_SECONDS = 5.0


class WordTemplateConfigError(OSError):
    """Base class for durable Word template configuration failures."""


class WordTemplateConfigCorruptError(WordTemplateConfigError):
    """Raised when an existing config cannot be decoded without data loss."""


class WordTemplateConfigConflictError(WordTemplateConfigError):
    """Raised when a compare-and-swap write uses a stale revision."""

    def __init__(self, current_revision):
        self.current_revision = int(current_revision)
        super().__init__("Cấu hình biểu mẫu Word đã được thay đổi.")


class WordTemplateConfigLockTimeoutError(WordTemplateConfigError):
    """Raised when another process holds the config mutation lock too long."""


def number_to_vietnamese_words(n):
    if n is None:
        return ""
    try:
        n = int(float(n))
    except (ValueError, TypeError):
        return ""
    if n == 0:
        return "Không"

    def read_three_digits(num, is_first=False):
        digits = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
        hundreds = num // 100
        tens = (num % 100) // 10
        ones = num % 10

        res = []
        if hundreds > 0 or not is_first:
            res.append(digits[hundreds])
            res.append("trăm")

        if tens > 0:
            if tens == 1:
                res.append("mười")
            else:
                res.append(digits[tens])
                res.append("mươi")
        elif (hundreds > 0 or not is_first) and ones > 0:
            res.append("lẻ")

        if ones > 0:
            if ones == 1 and tens > 1:
                res.append("mốt")
            elif ones == 5 and tens > 0:
                res.append("lăm")
            elif ones == 4 and tens > 1:
                res.append("tư")
            else:
                res.append(digits[ones])
        return " ".join(res)

    def read_block_9(num, is_first=False):
        million = num // 1000000
        thousand = (num % 1000000) // 1000
        unit = num % 1000

        res = []
        if million > 0:
            res.append(read_three_digits(million, is_first))
            res.append("triệu")
            is_first = False

        if thousand > 0:
            if million > 0 and res:
                res[-1] = res[-1] + ","
            res.append(read_three_digits(thousand, is_first))
            res.append("nghìn")
            is_first = False
        elif not is_first and unit > 0:
            if million > 0 and res:
                res[-1] = res[-1] + ","
            res.append("không trăm")

        if unit > 0:
            if thousand > 0 or (million > 0 and not is_first):
                if res:
                    res[-1] = res[-1] + ","
            res.append(read_three_digits(unit, is_first))

        return " ".join(res)

    billion_part = n // 1000000000
    rem_part = n % 1000000000

    parts = []
    if billion_part > 0:
        parts.append(number_to_vietnamese_words(billion_part))
        parts.append("tỷ,")
        if rem_part > 0:
            parts.append(read_block_9(rem_part, False))
        else:
            parts[-1] = "tỷ"
    else:
        parts.append(read_block_9(rem_part, True))

    words = " ".join(parts).strip()
    words = re.sub(r'\s+', ' ', words)
    words = words.replace(" ,", ",")
    if words:
        words = words[0].upper() + words[1:]
    return words

def enrich_context_with_lowercase_keys(data):
    if isinstance(data, dict):
        new_items = {}
        for k, v in list(data.items()):
            clean_k = k
            while clean_k.startswith('{') and clean_k.endswith('}'):
                clean_k = clean_k[1:-1].strip()

            enrich_context_with_lowercase_keys(v)
            new_items[clean_k.lower()] = v
            if clean_k != k:
                new_items[clean_k] = v
        data.update(new_items)
    elif isinstance(data, list):
        for item in data:
            enrich_context_with_lowercase_keys(item)

def enrich_context_with_words(data):
    if isinstance(data, dict):
        new_items = {}
        for k, v in list(data.items()):

            enrich_context_with_words(v)

            is_num = False
            num_val = None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                is_num = True
                num_val = v

            if is_num:
                words = number_to_vietnamese_words(num_val) + " đồng"

                clean_k = k
                while clean_k.startswith('{') and clean_k.endswith('}'):
                    clean_k = clean_k[1:-1].strip()


                new_items["bangchu_" + clean_k] = words
                new_items["BangChu_" + clean_k] = words
                new_items["bangchu_" + clean_k.lower()] = words
                new_items["BangChu_" + clean_k.lower()] = words
                new_items["bangchu_" + clean_k.upper()] = words
                new_items["BangChu_" + clean_k.upper()] = words


                s1 = re.sub('(.)([A-Z][a-z]+)', r'\1_\2', clean_k)
                snake_k = re.sub('([a-z0-9])([A-Z])', r'\1_\2', s1).lower()
                if snake_k != clean_k:
                    new_items["bangchu_" + snake_k] = words
                    new_items["BangChu_" + snake_k] = words
        data.update(new_items)
    elif isinstance(data, list):
        for item in data:
            enrich_context_with_words(item)

def normalize_date_str(val_str):
    if not isinstance(val_str, str):
        return val_str
    val_str = val_str.strip()


    m_iso = re.match(r'^(\d{4})-(\d{2})-(\d{2})T(\d{2}):(\d{2})(?::\d{2})?(?:\.\d+)?Z?$', val_str)
    if m_iso:
        y, m, d, hh, mm = m_iso.groups()
        return f"{d}/{m}/{y} {hh}:{mm}"


    m1 = re.match(r'^(\d{4})-(\d{2})-(\d{2})\s+(\d{2}):(\d{2})(?::\d{2})?$', val_str)
    if m1:
        y, m, d, hh, mm = m1.groups()
        return f"{d}/{m}/{y} {hh}:{mm}"


    m2 = re.match(r'^(\d{2})-(\d{2})-(\d{4})\s+(\d{2}):(\d{2})(?::\d{2})?$', val_str)
    if m2:
        d, m, y, hh, mm = m2.groups()
        return f"{d}/{m}/{y} {hh}:{mm}"


    m3 = re.match(r'^(\d{4})-(\d{2})-(\d{2})$', val_str)
    if m3:
        y, m, d = m3.groups()
        return f"{d}/{m}/{y}"


    m4 = re.match(r'^(\d{2})-(\d{2})-(\d{4})$', val_str)
    if m4:
        d, m, y = m4.groups()
        return f"{d}/{m}/{y}"

    return val_str


DATETIME_FIELD_NAMES = {
    'thoi_gian_dang_tai', 'thoi_gian_dang_ma', 'thoi_gian_dong_thau',
    'thoi_gian_mo_thau', 'thoi_gian_mo_ehsdxtc', 'thoi_gian_yeu_cau',
    'thoi_gian_tra_loi', 'thoi_gian', 'current_time',
}


def is_datetime_field_name(key_name):
    key = str(key_name or '').strip().strip('{}')
    key = re.sub(r'(?<!^)(?=[A-Z])', '_', key).lower()
    return key in DATETIME_FIELD_NAMES or any(
        key.endswith(suffix)
        for suffix in (
            'thoi_gian_dang_tai', 'thoi_gian_dang_ma', 'thoi_gian_dong_thau',
            'thoi_gian_mo_thau', 'thoi_gian_mo_ehsdxtc', 'thoi_gian_yeu_cau',
            'thoi_gian_tra_loi',
        )
    )

def format_vietnamese_datetime(val_str, key_name=None):
    if not isinstance(val_str, str):
        return val_str
    val_str = normalize_date_str(val_str)


    dt_match = re.match(r'^(\d{2})/(\d{2})/(\d{4})\s+(\d{2}):(\d{2})$', val_str)
    if dt_match:
        d, m, y, hh, mm = dt_match.groups()
        m_int = int(m)
        m_str = f"{m_int:02d}" if m_int in [1, 2] else str(m_int)
        if is_datetime_field_name(key_name):
            return f"{hh}:{mm} ngày {d}/{m_str}/{y}"
        return f"ngày {d} tháng {m_str} năm {y}"


    d_match = re.match(r'^(\d{2})/(\d{2})/(\d{4})$', val_str)
    if d_match:
        d, m, y = d_match.groups()
        m_int = int(m)
        if m_int in [1, 2]:
            m_str = f"{m_int:02d}"
        else:
            m_str = str(m_int)
        return f"ngày {d} tháng {m_str} năm {y}"

    return val_str
class SmartDate(str):
    def __sub__(self, other):
        try:
            def parse_to_datetime(s):
                if not s:
                    return None
                s = str(s).strip()

                m_speech = re.search(r'ngày\s+(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(\d{4})', s)
                if m_speech:
                    d, m, y = m_speech.groups()
                    return datetime(int(y), int(m), int(d))

                m_speech_t = re.search(r'(\d{1,2})\s+giờ\s+(\d{1,2})\s+phút\s+ngày\s+(\d{1,2})/(\d{1,2})/(\d{4})', s)
                if m_speech_t:
                    hh, mm, d, m, y = m_speech_t.groups()
                    return datetime(int(y), int(m), int(d), int(hh), int(mm))

                m_compact_t = re.search(r'(\d{1,2}):(\d{2})\s+ngày\s+(\d{1,2})/(\d{1,2})/(\d{4})', s)
                if m_compact_t:
                    hh, mm, d, m, y = m_compact_t.groups()
                    return datetime(int(y), int(m), int(d), int(hh), int(mm))

                m_dt = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})\s+(\d{1,2}):(\d{1,2})', s)
                if m_dt:
                    d, m, y, hh, mm = m_dt.groups()
                    return datetime(int(y), int(m), int(d), int(hh), int(mm))

                m_d = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', s)
                if m_d:
                    d, m, y = m_d.groups()
                    return datetime(int(y), int(m), int(d))

                for fmt in ["%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"]:
                    try:
                        return datetime.strptime(s, fmt)
                    except ValueError:
                        continue
                return None

            dt_self = parse_to_datetime(self)
            dt_other = parse_to_datetime(other)
            if dt_self and dt_other:
                return (dt_self - dt_other).days
        except (TypeError, ValueError, OverflowError):
            pass
        return ''

    def __rsub__(self, other):
        try:


            return SmartDate(other).__sub__(self)
        except (TypeError, ValueError, OverflowError):
            pass
        return ''

def format_context_dates(data):
    if isinstance(data, dict):
        new_items = {}
        for k, v in list(data.items()):

            format_context_dates(v)

            if isinstance(v, str):
                is_date_key = any(x in k.lower() for x in ['ngay', 'thoi_gian', 'date', 'time', 'mo_thau', 'dong_thau', 'dang_tai', 'ky'])
                v_norm = normalize_date_str(v)
                is_date = False
                date_only_val = None
                year_val = None


                dt_match = re.match(r'^(\d{2})/(\d{2})/(\d{4})\s+(\d{2}):(\d{2})$', v_norm)
                if dt_match:
                    is_date = True
                    d, m, y = dt_match.group(1), dt_match.group(2), dt_match.group(3)
                    m_int = int(m)
                    m_str = f"{m_int:02d}" if m_int in [1, 2] else str(m_int)
                    date_only_val = f"{d}/{m_str}/{y}"
                    year_val = y


                d_match = re.match(r'^(\d{2})/(\d{2})/(\d{4})$', v_norm)
                if d_match:
                    is_date = True
                    d, m, y = d_match.group(1), d_match.group(2), d_match.group(3)
                    m_int = int(m)
                    m_str = f"{m_int:02d}" if m_int in [1, 2] else str(m_int)
                    date_only_val = f"{d}/{m_str}/{y}"
                    year_val = y

                if is_date:

                    data[k] = SmartDate(format_vietnamese_datetime(v_norm, key_name=k))


                    clean_k = k
                    while clean_k.startswith('{') and clean_k.endswith('}'):
                        clean_k = clean_k[1:-1].strip()

                    new_items["S_" + clean_k] = SmartDate(date_only_val)
                    new_items["s_" + clean_k] = SmartDate(date_only_val)
                    new_items["S_" + clean_k.lower()] = SmartDate(date_only_val)
                    new_items["s_" + clean_k.lower()] = SmartDate(date_only_val)
                    new_items["S_" + clean_k.upper()] = SmartDate(date_only_val)
                    new_items["s_" + clean_k.upper()] = SmartDate(date_only_val)


                    new_items["nam_" + clean_k] = SmartDate(year_val)
                    new_items["Nam_" + clean_k] = SmartDate(year_val)
                    new_items["NAM_" + clean_k] = SmartDate(year_val)
                    new_items["nam_" + clean_k.lower()] = SmartDate(year_val)
                    new_items["Nam_" + clean_k.lower()] = SmartDate(year_val)
                    new_items["NAM_" + clean_k.lower()] = SmartDate(year_val)
                    new_items["nam_" + clean_k.upper()] = SmartDate(year_val)
                    new_items["Nam_" + clean_k.upper()] = SmartDate(year_val)
                    new_items["NAM_" + clean_k.upper()] = SmartDate(year_val)


                    s1 = re.sub('(.)([A-Z][a-z]+)', r'\1_\2', clean_k)
                    snake_k = re.sub('([a-z0-9])([A-Z])', r'\1_\2', s1).lower()
                    if snake_k != clean_k:
                        new_items["s_" + snake_k] = SmartDate(date_only_val)
                        new_items["nam_" + snake_k] = SmartDate(year_val)
                elif is_date_key:
                    data[k] = SmartDate(v)
        data.update(new_items)
    elif isinstance(data, list):
        for item in data:
            format_context_dates(item)
project_root = str(PROJECT_ROOT)
TEMPLATE_DIR = str(WORD_TEMPLATE_DIR)
DEFAULT_TEMPLATE = ''
LEGACY_WORD_TEMPLATES = {
    'mau_bao_cao_dau_thau.docx': 'Bản báo cáo đánh giá mặc định',
    'mau_hop_dong_lcnt.docx': 'Mẫu hợp đồng kinh tế LCNT',
}

def get_user_template_dir(user_id=None, *, create=True):
    if user_id:
        clean_user_id = str(user_id).replace('..', '').replace('/', '').replace('\\', '').strip()
        path = os.path.join(TEMPLATE_DIR, clean_user_id)
        if create:
            os.makedirs(path, exist_ok=True)
        return path
    return TEMPLATE_DIR

def get_scope_template_dir(owner_type, owner_id, *, create=True):
    normalized_type = str(owner_type or '').strip().lower()
    normalized_id = str(owner_id or '').replace('..', '').replace('/', '').replace('\\', '').strip()
    if not normalized_id:
        raise ValueError('Phạm vi biểu mẫu Word không hợp lệ')
    if normalized_type == 'personal':
        return get_user_template_dir(normalized_id, create=create)
    if normalized_type != 'organization':
        raise ValueError('Loại phạm vi biểu mẫu Word không hợp lệ')
    organizations_dir = os.path.realpath(os.path.join(TEMPLATE_DIR, 'organizations'))
    if create:
        os.makedirs(organizations_dir, exist_ok=True)
    path = os.path.realpath(os.path.join(organizations_dir, normalized_id))
    if not path.startswith(organizations_dir + os.sep):
        raise ValueError('Phạm vi biểu mẫu Word không hợp lệ')
    if create:
        os.makedirs(path, exist_ok=True)
    return path

def _template_config_path(owner_id=None, *, owner_type='personal', create=False):
    scope_dir = (
        get_scope_template_dir(owner_type, owner_id, create=create)
        if owner_id else TEMPLATE_DIR
    )
    return os.path.join(scope_dir, 'config.json')


def _template_config_revision(config):
    revision = config.get('revision', 0)
    if isinstance(revision, bool) or not isinstance(revision, int) or revision < 0:
        raise WordTemplateConfigCorruptError(
            'Revision cấu hình biểu mẫu Word không hợp lệ.'
        )
    return revision


def _load_template_config(owner_id=None, *, owner_type='personal'):
    config_path = _template_config_path(owner_id, owner_type=owner_type)
    if not os.path.exists(config_path):
        return {}
    try:
        with open(config_path, 'r', encoding='utf-8') as file_obj:
            config = json.load(file_obj)
    except (json.JSONDecodeError, UnicodeDecodeError, TypeError) as error:
        raise WordTemplateConfigCorruptError(
            'Cấu hình biểu mẫu Word bị hỏng; dữ liệu gốc được giữ nguyên.'
        ) from error
    if not isinstance(config, dict):
        raise WordTemplateConfigCorruptError(
            'Cấu hình biểu mẫu Word phải là một JSON object.'
        )
    _template_config_revision(config)
    return config


@contextmanager
def _exclusive_file_lock(
    lock_path,
    *,
    timeout_seconds=_WORD_CONFIG_LOCK_TIMEOUT_SECONDS,
):
    deadline = time.monotonic() + max(float(timeout_seconds), 0.0)
    with open(lock_path, 'a+b') as lock_file:
        if os.name == 'nt':
            import msvcrt

            lock_file.seek(0, os.SEEK_END)
            if lock_file.tell() == 0:
                lock_file.write(b'\0')
                lock_file.flush()
            while True:
                try:
                    lock_file.seek(0)
                    msvcrt.locking(lock_file.fileno(), msvcrt.LK_NBLCK, 1)
                    break
                except OSError as error:
                    if time.monotonic() >= deadline:
                        raise WordTemplateConfigLockTimeoutError(
                            'Cấu hình biểu mẫu Word đang được cập nhật.'
                        ) from error
                    time.sleep(0.01)
            try:
                yield
            finally:
                lock_file.seek(0)
                msvcrt.locking(lock_file.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl

            while True:
                try:
                    fcntl.flock(
                        lock_file.fileno(),
                        fcntl.LOCK_EX | fcntl.LOCK_NB,
                    )
                    break
                except BlockingIOError as error:
                    if time.monotonic() >= deadline:
                        raise WordTemplateConfigLockTimeoutError(
                            'Cấu hình biểu mẫu Word đang được cập nhật.'
                        ) from error
                    time.sleep(0.01)
            try:
                yield
            finally:
                fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)


@contextmanager
def _template_config_file_lock(
    owner_id=None,
    *,
    owner_type='personal',
    timeout_seconds=_WORD_CONFIG_LOCK_TIMEOUT_SECONDS,
):
    config_path = _template_config_path(
        owner_id,
        owner_type=owner_type,
        create=True,
    )
    with _exclusive_file_lock(
        f'{config_path}.lock',
        timeout_seconds=timeout_seconds,
    ):
        yield


@contextmanager
def template_scope_file_lock(
    owner_id,
    *,
    owner_type='personal',
    timeout_seconds=_WORD_CONFIG_LOCK_TIMEOUT_SECONDS,
):
    scope_dir = get_scope_template_dir(owner_type, owner_id, create=True)
    with _exclusive_file_lock(
        os.path.join(scope_dir, '.templates.lock'),
        timeout_seconds=timeout_seconds,
    ):
        yield


def _write_template_config(config, owner_id=None, *, owner_type='personal'):
    config_path = _template_config_path(
        owner_id,
        owner_type=owner_type,
        create=True,
    )
    temporary_path = f"{config_path}.{uuid.uuid4().hex}.tmp"
    try:
        with open(temporary_path, 'x', encoding='utf-8') as file_obj:
            json.dump(config, file_obj, ensure_ascii=False, indent=4, sort_keys=True)
            file_obj.flush()
            os.fsync(file_obj.fileno())
        os.replace(temporary_path, config_path)
    finally:
        if os.path.exists(temporary_path):
            os.remove(temporary_path)


def _mutate_template_config(
    mutation,
    owner_id=None,
    *,
    owner_type='personal',
    expected_revision=None,
    commit_callback=None,
):
    with _WORD_CONFIG_LOCK:
        with _template_config_file_lock(owner_id, owner_type=owner_type):
            config = _load_template_config(owner_id, owner_type=owner_type)
            original_config = copy.deepcopy(config)
            config_path = _template_config_path(
                owner_id,
                owner_type=owner_type,
            )
            config_existed = os.path.exists(config_path)
            current_revision = _template_config_revision(config)
            if (
                expected_revision is not None
                and expected_revision != current_revision
            ):
                raise WordTemplateConfigConflictError(current_revision)
            changed = mutation(config)
            if changed is False:
                if commit_callback is not None:
                    commit_callback(current_revision)
                return current_revision
            next_revision = current_revision + 1
            config['revision'] = next_revision
            _write_template_config(config, owner_id, owner_type=owner_type)
            if commit_callback is not None:
                try:
                    commit_callback(next_revision)
                except Exception:
                    if config_existed:
                        _write_template_config(
                            original_config,
                            owner_id,
                            owner_type=owner_type,
                        )
                    elif os.path.exists(config_path):
                        os.remove(config_path)
                    raise
            return next_revision


def get_template_config_revision(owner_id=None, *, owner_type='personal'):
    with _WORD_CONFIG_LOCK:
        config = _load_template_config(owner_id, owner_type=owner_type)
        return _template_config_revision(config)


def get_active_template(owner_id=None, *, owner_type='personal'):
    with _WORD_CONFIG_LOCK:
        config = _load_template_config(owner_id, owner_type=owner_type)
        return str(config.get('active_template') or DEFAULT_TEMPLATE)


def _enabled_template_names_from_config(config):
    if 'enabled_templates' in config:
        return _normalize_template_assignment_values(
            config.get('enabled_templates')
        )
    legacy_names = []
    active_template = str(config.get('active_template') or '').strip()
    if active_template:
        legacy_names.append(active_template)
    assignments = config.get('template_assignments')
    if isinstance(assignments, dict):
        for filenames in assignments.values():
            legacy_names.extend(_normalize_template_assignment_values(filenames))
    return _normalize_template_assignment_values(legacy_names)


def get_enabled_templates(owner_id=None, *, owner_type='personal'):
    with _WORD_CONFIG_LOCK:
        config = _load_template_config(owner_id, owner_type=owner_type)
        return _enabled_template_names_from_config(config)


def is_template_enabled(filename, owner_id=None, *, owner_type='personal'):
    identity = str(filename or '').strip().casefold()
    if not identity:
        return False
    return any(
        item.casefold() == identity
        for item in get_enabled_templates(owner_id, owner_type=owner_type)
    )


def configure_template_availability(
    filename,
    enabled,
    owner_id=None,
    *,
    owner_type='personal',
    expected_revision=None,
    activate=False,
    commit_callback=None,
):
    safe_name = str(filename or '').strip()
    if not safe_name:
        raise ValueError('Tên biểu mẫu không được để trống')
    enabled_templates = []

    def mutate(config):
        nonlocal enabled_templates
        enabled_templates = _enabled_template_names_from_config(config)
        target_identity = safe_name.casefold()
        enabled_templates = [
            item for item in enabled_templates
            if item.casefold() != target_identity
        ]
        if enabled:
            enabled_templates.append(safe_name)
        config['enabled_templates'] = enabled_templates

        if activate and enabled:
            config['active_template'] = safe_name
        active_template = str(config.get('active_template') or '').strip()
        enabled_identities = {
            item.casefold() for item in enabled_templates
        }
        if active_template.casefold() not in enabled_identities:
            config['active_template'] = (
                enabled_templates[0] if enabled_templates else ''
            )
    revision = _mutate_template_config(
        mutate,
        owner_id,
        owner_type=owner_type,
        expected_revision=expected_revision,
        commit_callback=commit_callback,
    )
    return enabled_templates, revision


def set_template_enabled(
    filename,
    enabled,
    owner_id=None,
    *,
    owner_type='personal',
    expected_revision=None,
):
    enabled_templates, _revision = configure_template_availability(
        filename,
        enabled,
        owner_id,
        owner_type=owner_type,
        expected_revision=expected_revision,
    )
    return enabled_templates


def get_template_assignments(owner_id=None, *, owner_type='personal'):
    with _WORD_CONFIG_LOCK:
        config = _load_template_config(owner_id, owner_type=owner_type)
        assignments = config.get('template_assignments')
        if not isinstance(assignments, dict):
            return {}
        normalized = {}
        for document_type, filenames in assignments.items():
            safe_type = str(document_type).strip()
            safe_filenames = _normalize_template_assignment_values(filenames)
            if safe_type and safe_filenames:
                normalized[safe_type] = safe_filenames
        return normalized


def _normalize_template_assignment_values(value):
    candidates = [value] if isinstance(value, str) else value
    if not isinstance(candidates, (list, tuple)):
        return []
    normalized = []
    seen = set()
    for filename in candidates:
        if not isinstance(filename, str):
            continue
        safe_name = filename.strip()
        identity = safe_name.casefold()
        if not safe_name or identity in seen:
            continue
        seen.add(identity)
        normalized.append(safe_name)
    return normalized


def set_template_assignments(
    assignments,
    owner_id=None,
    *,
    owner_type='personal',
    expected_revision=None,
    commit_callback=None,
):
    normalized = {}
    for document_type, filenames in dict(assignments or {}).items():
        safe_type = str(document_type).strip()
        safe_filenames = _normalize_template_assignment_values(filenames)
        if safe_type and safe_filenames:
            normalized[safe_type] = safe_filenames
    def mutate(config):
        config['template_assignments'] = normalized
    _mutate_template_config(
        mutate,
        owner_id,
        owner_type=owner_type,
        expected_revision=expected_revision,
        commit_callback=commit_callback,
    )
    return normalized


def replace_template_reference(
    current_filename,
    next_filename,
    owner_id=None,
    *,
    owner_type='personal',
    commit_callback=None,
):
    def mutate(config):
        changed = False
        if config.get('active_template') == current_filename:
            config['active_template'] = next_filename
            changed = True
        enabled_templates = config.get('enabled_templates')
        if isinstance(enabled_templates, (list, tuple)):
            values = _normalize_template_assignment_values(enabled_templates)
            if current_filename in values:
                updated = [
                    next_filename if filename == current_filename else filename
                    for filename in values
                    if filename != current_filename or next_filename
                ]
                config['enabled_templates'] = (
                    _normalize_template_assignment_values(updated)
                )
                changed = True
        assignments = config.get('template_assignments')
        if isinstance(assignments, dict):
            for document_type, filenames in list(assignments.items()):
                values = _normalize_template_assignment_values(filenames)
                if current_filename not in values:
                    continue
                updated = [
                    next_filename if filename == current_filename else filename
                    for filename in values
                    if filename != current_filename or next_filename
                ]
                updated = _normalize_template_assignment_values(updated)
                if updated:
                    assignments[document_type] = updated
                else:
                    assignments.pop(document_type, None)
                changed = True
        return changed

    return _mutate_template_config(
        mutate,
        owner_id,
        owner_type=owner_type,
        commit_callback=commit_callback,
    )


def resolve_publication_templates(
    document_type,
    owner_id=None,
    *,
    owner_type='personal',
    allow_active_fallback=False,
):
    assignments = get_template_assignments(owner_id, owner_type=owner_type)
    assigned = assignments.get(str(document_type or '').strip(), [])
    if assigned:
        enabled = {
            filename.casefold()
            for filename in get_enabled_templates(
                owner_id,
                owner_type=owner_type,
            )
        }
        return [
            filename for filename in assigned
            if filename.casefold() in enabled
        ], 'assignment'
    if allow_active_fallback:
        active = get_active_template(owner_id, owner_type=owner_type)
        if active and is_template_enabled(
            active,
            owner_id,
            owner_type=owner_type,
        ):
            return [active], 'legacy-active'
    return [], 'unassigned'


def resolve_publication_template(
    document_type,
    owner_id=None,
    *,
    owner_type='personal',
    allow_active_fallback=False,
):
    templates, source = resolve_publication_templates(
        document_type,
        owner_id,
        owner_type=owner_type,
        allow_active_fallback=allow_active_fallback,
    )
    return (templates[0] if templates else ''), source


def set_active_template(
    filename,
    owner_id=None,
    *,
    owner_type='personal',
    expected_revision=None,
):
    def mutate(config):
        config['active_template'] = filename
    _mutate_template_config(
        mutate,
        owner_id,
        owner_type=owner_type,
        expected_revision=expected_revision,
    )

def list_templates(owner_id=None, *, owner_type='personal'):
    templates = []
    active_template = get_active_template(owner_id, owner_type=owner_type)
    enabled_templates = {
        filename.casefold()
        for filename in get_enabled_templates(owner_id, owner_type=owner_type)
    }
    scope_dir = get_scope_template_dir(owner_type, owner_id, create=False) if owner_id else TEMPLATE_DIR
    scoped_filenames = set()
    if os.path.exists(scope_dir):
        for f in sorted(os.listdir(scope_dir), key=str.casefold):
            if f.lower().endswith('.docx'):
                if not owner_id and f.lower() in LEGACY_WORD_TEMPLATES:
                    continue
                scoped_filenames.add(f.lower())
                templates.append({
                    'filename': f,
                    'name': f,
                    'is_system': False,
                    'is_mutable': True,
                    'is_available': True,
                    'is_enabled': f.casefold() in enabled_templates,
                    'is_active': active_template == f
                })

    # Older installations may still contain the two templates that used to be
    # provisioned globally. Keep them visible and manageable until an operator
    # edits or deletes them; fresh installations no longer create these files.
    for filename, name in LEGACY_WORD_TEMPLATES.items():
        if filename.lower() in scoped_filenames:
            continue
        if os.path.isfile(os.path.join(TEMPLATE_DIR, filename)):
            templates.append({
                'filename': filename,
                'name': name,
                'is_system': False,
                'is_mutable': True,
                'is_legacy': True,
                'is_available': True,
                'is_enabled': filename.casefold() in enabled_templates,
                'is_active': active_template == filename,
            })
    return templates

def translate_xml_tags(xml_content, valid_vars):


    def clean_braces(text):
        result = []
        i = 0
        n = len(text)
        while i < n:
            if text[i] == '{':
                j = i + 1
                brace_count = 1
                while j < n:
                    if text[j] == '{':
                        brace_count += 1
                    elif text[j] == '}':
                        brace_count -= 1
                        if brace_count == 0:
                            break
                    j += 1
                if j < n:
                    inside = text[i+1:j]
                    inside_clean = re.sub(r'<[^>]*>', '', inside)
                    result.append('{' + inside_clean + '}')
                    i = j + 1
                else:
                    result.append(text[i])
                    i += 1
            else:
                result.append(text[i])
                i += 1
        return "".join(result)

    xml_content = clean_braces(xml_content)


    xml_content = re.sub(r'\{%\s+p\s+', r'{%p ', xml_content)
    xml_content = re.sub(r'\{%\s+tr\s+', r'{%tr ', xml_content)
    xml_content = re.sub(r'\{%\s+tc\s+', r'{%tc ', xml_content)


    def replace_generic_loop(match):
        loop_name = match.group(1).lower()
        block_content = match.group(2)

        def replace_var(var_match):
            raw_var = var_match.group(1)
            var_name = raw_var.lower()
            if var_name.startswith('item.'):
                var_name = var_name[5:]

            if var_name.startswith('#') or var_name.startswith('/') or var_name.startswith('%') or var_name.startswith('^'):
                return var_match.group(0)
            if var_name in valid_vars:
                return f"{{{{ item.{var_name} }}}}"
            return var_match.group(0)

        new_content = re.sub(r'(?<!\{)\{((?:item\.)?[A-Za-z0-9_]+)\}(?!\})', replace_var, block_content, flags=re.IGNORECASE)
        return f"{{#{loop_name}}}{new_content}{{/{loop_name}}}"

    xml_content = re.sub(r'\{#([A-Za-z0-9_]+)\}(.*?)\{/\1\}', replace_generic_loop, xml_content, flags=re.DOTALL)



    def replace_open_loop(match):
        loop_name = match.group(1).lower()
        if loop_name not in valid_vars:
            return match.group(0)
        index = match.start()


        pos_start = xml_content.rfind('<w:tr', 0, index)
        pos_end = xml_content.rfind('</w:tr>', 0, index)
        in_table_row = pos_start > pos_end

        if in_table_row:
            return f'{{%tr for item in {loop_name} %}}'
        else:
            return f'{{% for item in {loop_name} %}}'

    def replace_close_loop(match):
        loop_name = match.group(1).lower()
        if loop_name not in valid_vars:
            return match.group(0)
        index = match.start()

        pos_start = xml_content.rfind('<w:tr', 0, index)
        pos_end = xml_content.rfind('</w:tr>', 0, index)
        in_table_row = pos_start > pos_end

        if in_table_row:
            return '{%tr endfor %}'
        else:
            return '{% endfor %}'


    xml_content = re.sub(r'\{#([A-Za-z0-9_]+)\}', replace_open_loop, xml_content)
    xml_content = re.sub(r'\{/([A-Za-z0-9_]+)\}', replace_close_loop, xml_content)


    def replace_global_var(match):
        var_name = match.group(1).lower()
        if var_name.startswith('#') or var_name.startswith('/') or var_name.startswith('%') or var_name.startswith('^'):
            return match.group(0)
        if var_name in valid_vars:
            return f"{{{{ {var_name} }}}}"
        return match.group(0)

    xml_content = re.sub(r'(?<!\{)\{([A-Za-z0-9_]+)\}(?!\})', replace_global_var, xml_content)


    def pull_tr_loops_out(xml):
        def process_tr(match):
            tr_content = match.group(0)
            start_matches = list(re.finditer(r'\{%\s*tr\s+(for\s+.*?)\s*%}', tr_content))
            end_matches = list(re.finditer(r'\{%\s*tr\s+endfor\s*%}', tr_content))

            if start_matches and end_matches:
                start_tag = start_matches[0].group(0)
                loop_expr = start_matches[0].group(1)
                end_tag = end_matches[-1].group(0)

                cleaned_tr = tr_content.replace(start_tag, '').replace(end_tag, '')
                return f"{{% {loop_expr} %}}{cleaned_tr}{{% endfor %}}"
            return tr_content

        return re.sub(r'<w:tr[ >].*?</w:tr>', process_tr, xml, flags=re.DOTALL)

    xml_content = pull_tr_loops_out(xml_content)
    return xml_content

_TRANSLATED_DOCXTPL_CACHE = {}

def translate_docx_template(template_path, context, allowed_root_keys=None):

    global _TRANSLATED_DOCXTPL_CACHE

    mtime = os.path.getmtime(template_path)

    context.update(COLUMN_LITERAL_CONTEXT)
    if allowed_root_keys is not None:
        allowed_root_keys = set(allowed_root_keys) | set(COLUMN_LITERAL_CONTEXT)
    enrich_context_with_lowercase_keys(context)
    valid_vars = set(context.keys())
    for val in context.values():
        if isinstance(val, list):
            for item in val:
                if isinstance(item, dict):
                    valid_vars.update(item.keys())
        elif isinstance(val, dict):
            valid_vars.update(val.keys())

    valid_vars_hash = hash(frozenset(valid_vars))

    cached = _TRANSLATED_DOCXTPL_CACHE.get(template_path)
    if cached and cached[0] == mtime and cached[1] == valid_vars_hash:

        return DocxTemplate(BytesIO(cached[3]))


    with open(template_path, 'rb') as template_file:
        template_bytes = template_file.read()
    validate_ooxml_archive(template_bytes, "docx")

    expanded_template_bytes = expand_column_loops(template_bytes, context)
    has_column_loops = expanded_template_bytes is not template_bytes
    template_bytes = expanded_template_bytes

    temp_bytes = BytesIO()
    translated_xml_parts = []
    with zipfile.ZipFile(BytesIO(template_bytes), 'r') as yin:
        with zipfile.ZipFile(temp_bytes, 'w', zipfile.ZIP_DEFLATED) as yout:
            for item in yin.infolist():
                data = yin.read(item.filename)
                if item.filename in ['word/document.xml', 'word/header1.xml', 'word/header2.xml', 'word/footer1.xml', 'word/footer2.xml']:
                    xml_str = data.decode('utf-8')
                    translated_xml = translate_xml_tags(xml_str, valid_vars)
                    data = translated_xml.encode('utf-8')
                if item.filename.startswith('word/') and item.filename.lower().endswith('.xml'):
                    translated_xml_parts.append(data.decode('utf-8'))
                yout.writestr(item, data)

    validate_template_statements(translated_xml_parts)
    if allowed_root_keys is not None:
        validate_template_root_keys(translated_xml_parts, allowed_root_keys)

    translated_data = temp_bytes.getvalue()

    if not has_column_loops:
        _TRANSLATED_DOCXTPL_CACHE[template_path] = (
            mtime,
            valid_vars_hash,
            DocxTemplate(BytesIO(translated_data)),
            translated_data,
        )

    return DocxTemplate(BytesIO(translated_data))

def replace_placeholders_with_empty(data):
    if isinstance(data, dict):
        return {k: replace_placeholders_with_empty(v) for k, v in data.items()}
    elif isinstance(data, list):
        return [replace_placeholders_with_empty(x) for x in data]
    elif data == '--':
        return ''
    return data

_OPTIMIZED_IMAGE_CACHE = {}

def optimize_image_for_docx(filepath, max_width=800):
    try:
        mtime = os.path.getmtime(filepath)
        cache_key = (filepath, max_width)


        if cache_key in _OPTIMIZED_IMAGE_CACHE:
            cached_mtime, cached_data = _OPTIMIZED_IMAGE_CACHE[cache_key]
            if cached_mtime == mtime:
                return BytesIO(cached_data)


        dir_name = os.path.dirname(filepath)
        base_name = os.path.basename(filepath)
        name, ext = os.path.splitext(base_name)
        cache_filename = f"{name}_opt_{max_width}.jpg"
        cache_path = os.path.join(dir_name, cache_filename)


        if os.path.exists(cache_path) and os.path.getmtime(cache_path) >= mtime:
            with open(cache_path, "rb") as f:
                data = f.read()
            _OPTIMIZED_IMAGE_CACHE[cache_key] = (mtime, data)
            return BytesIO(data)



        file_size = os.path.getsize(filepath)
        if ext.lower() in ['.jpg', '.jpeg'] and file_size < 50000:
            with open(filepath, "rb") as f:
                data = f.read()
            _OPTIMIZED_IMAGE_CACHE[cache_key] = (mtime, data)
            return BytesIO(data)

        from PIL import Image
        with Image.open(filepath) as img:
            w, h = img.size
            if w > max_width:
                ratio = max_width / w
                new_w = int(w * ratio)
                new_h = int(h * ratio)

                resample = Image.Resampling.BOX if max_width <= 300 else Image.Resampling.BILINEAR
                img = img.resize((new_w, new_h), resample)


            if img.mode in ('RGBA', 'LA', 'P'):
                background = Image.new("RGB", img.size, (255, 255, 255))
                if img.mode == 'RGBA':
                    background.paste(img, mask=img.split()[3])
                else:
                    background.paste(img)
                img = background
            elif img.mode != 'RGB':
                img = img.convert('RGB')

            out = BytesIO()
            img.save(out, format="JPEG", quality=80)
            data = out.getvalue()


            try:
                with open(cache_path, "wb") as f:
                    f.write(data)
            except OSError as cache_err:
                log_error(cache_err, "Document.ImageCacheWrite")

            _OPTIMIZED_IMAGE_CACHE[cache_key] = (mtime, data)
            out.seek(0)
            return out
    except Exception as e:
        log_error(e, "Document.ImageOptimization")
        return filepath

def prewarm_image_cache():

    try:
        images_dir = os.path.join(IMAGE_DIR, 'chuyen_gia')
        if not os.path.exists(images_dir):
            return

        tasks = []
        for fname in os.listdir(images_dir):

            if '_opt_' in fname:
                continue
            if not any(fname.lower().endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.webp']):
                continue

            fpath = os.path.join(images_dir, fname)
            max_w = 1200 if '_cert' in fname else 300


            name_no_ext, _ = os.path.splitext(fname)
            cache_path = os.path.join(images_dir, f"{name_no_ext}_opt_{max_w}.jpg")
            if os.path.exists(cache_path) and os.path.getmtime(cache_path) >= os.path.getmtime(fpath):
                continue

            tasks.append((fpath, max_w))

        if not tasks:
            return


        futures = [_IMAGE_THREAD_POOL.submit(optimize_image_for_docx, fpath, max_w) for fpath, max_w in tasks]
        count = 0
        for f in as_completed(futures):
            try:
                f.result()
                count += 1
            except Exception as e:
                log_error(e, "Document.ImagePrewarm")

        if count > 0 and os.environ.get("APP_DEBUG", "False").lower() == "true":
            append_runtime_log("export_error.log", f"Image prewarm completed: {count} files.\n")
    except Exception as e:
        log_error(e, "Document.ImagePrewarm")


def _resolve_docx_image_path(
    value,
    expected_subfolder,
    media_organization_id=None,
):
    if not isinstance(value, str) or expected_subfolder not in {"chuyen_gia", "nha_thau"}:
        return "", ""
    normalized = normalize_managed_image_path(value)
    expected_prefix = f"images/{expected_subfolder}/"
    if not normalized.startswith(expected_prefix):
        return "", ""
    parts = normalized.split("/")
    if len(parts) == 4 and (
        not media_organization_id
        or not managed_image_path_matches_tenant(
            normalized,
            media_organization_id,
        )
    ):
        return "", ""
    if len(parts) not in {3, 4}:
        return "", ""
    filename = parts[-1]
    if os.path.splitext(filename)[1].lower() not in {".png", ".jpg", ".jpeg", ".webp"}:
        return "", ""

    images_root = os.path.realpath(IMAGE_DIR)
    allowed_root = os.path.realpath(os.path.join(images_root, expected_subfolder))
    filepath = os.path.realpath(
        os.path.join(allowed_root, *parts[2:])
    )
    try:
        if os.path.commonpath([allowed_root, filepath]) != allowed_root:
            return "", ""
    except ValueError:
        return "", ""
    if not os.path.isfile(filepath):
        return "", ""
    return normalized, filepath


def _collect_image_tasks(
    data,
    tasks=None,
    allowed_image_fields=None,
    media_organization_id=None,
):

    if tasks is None:
        tasks = []
    if allowed_image_fields is None:
        allowed_image_fields = BASE_IMAGE_FIELDS
    if isinstance(data, dict):
        for k, v in list(data.items()):
            if isinstance(v, str):
                expected_subfolder = allowed_image_fields.get(str(k))
                norm_v, filepath = _resolve_docx_image_path(
                    v,
                    expected_subfolder,
                    media_organization_id,
                )
                if filepath:
                    data[k] = norm_v
                    max_w = 300
                    width_hint = 'small'
                    if 'chung_chi' in k or 'cert' in k:
                        max_w = 1200
                        width_hint = 'full'
                    tasks.append((data, k, filepath, max_w, width_hint))
            else:
                _collect_image_tasks(
                    v,
                    tasks,
                    allowed_image_fields,
                    media_organization_id,
                )
    elif isinstance(data, list):
        for item in data:
            _collect_image_tasks(
                item,
                tasks,
                allowed_image_fields,
                media_organization_id,
            )
    return tasks


def convert_images_in_context(
    doc,
    data,
    allowed_image_fields=None,
    media_organization_id=None,
):


    try:
        section = doc.sections[0]
        usable_width = section.page_width - section.left_margin - section.right_margin
    except (AttributeError, IndexError, TypeError):
        usable_width = Inches(6.0)


    tasks = _collect_image_tasks(
        data,
        allowed_image_fields=(
            BASE_IMAGE_FIELDS
            if allowed_image_fields is None
            else allowed_image_fields
        ),
        media_organization_id=media_organization_id,
    )
    if not tasks:
        return


    def _process_one(task):
        data_ref, k, filepath, max_w, width_hint = task
        image_stream = optimize_image_for_docx(filepath, max_w)
        return data_ref, k, image_stream, width_hint

    # Rendering runs in a process-creation-denied sandbox. Process images
    # sequentially so the worker never needs to clone threads/processes.
    for task in tasks:
        try:
            data_ref, k, image_stream, width_hint = _process_one(task)
            width_val = usable_width if width_hint == 'full' else Inches(1.5)
            data_ref[k] = InlineImage(doc, image_stream, width=width_val)
        except Exception as img_ex:
            log_error(img_ex, "Document.ImageConversion")

class TemplateRenderError(ValueError):
    """Public, non-sensitive error raised when a DOCX template cannot render."""


def generate_report_from_custom_template(
    template_path,
    context,
    context_manifest=None,
):

    if context_manifest is not None:
        render_policy = validate_docx_context_manifest(context, context_manifest)
        allowed_root_keys = render_policy["allowed_root_keys"]
        allowed_image_fields = render_policy["allowed_image_fields"]
        media_organization_id = render_policy["media_organization_id"]
    else:
        allowed_root_keys = set(context or {})
        allowed_image_fields = BASE_IMAGE_FIELDS
        media_organization_id = None
    context = replace_placeholders_with_empty(context)

    enrich_context_with_words(context)

    format_context_dates(context)

    doc = None
    try:
        doc = translate_docx_template(
            template_path,
            context,
            allowed_root_keys=allowed_root_keys,
        )
        convert_images_in_context(
            doc,
            context,
            allowed_image_fields=allowed_image_fields,
            media_organization_id=media_organization_id,
        )
        doc.render(context, jinja_env=create_template_environment())
    except Exception as e:
        try:
            # Do not persist the document context: it may contain identities,
            # signatures, email addresses, access tokens, or embedded files.
            append_runtime_log(
                "export_error.log",
                (
                    f"[{datetime.now(timezone.utc).isoformat()}] ERROR: Failed rendering "
                    f"template {os.path.basename(template_path)}\n"
                    f"{traceback.format_exc()}\n{'=' * 50}\n"
                ),
            )
        except OSError as log_write_error:
            log_error(log_write_error, "Document.RenderErrorLog")


        raise TemplateRenderError(
            "Mẫu Word chứa biểu thức không được hỗ trợ hoặc không thể kết xuất."
        ) from e

    out_stream = BytesIO()
    if doc:
        doc.save(out_stream)
    out_stream.seek(0)
    return out_stream
