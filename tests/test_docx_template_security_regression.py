import io
import zipfile
from xml.etree import ElementTree

import pytest
from docx import Document
from docxtpl import DocxTemplate
from jinja2 import UndefinedError

from backend.documents.template_security import create_template_environment


def _template_bytes():
    document = Document()
    document.add_paragraph("{{ value }}")
    document.add_paragraph("{% for item in items %}{{ item }}{% endfor %}")
    document.add_paragraph("{% if enabled %}YES{% else %}NO{% endif %}")
    document.add_paragraph("{{ value|upper }}")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize(
    "payload",
    ["A & B", "<abc>", '\"</w:t><w:r><w:t>INJECTED', "{{ variable }}", "Tiếng Việt", "line1\nline2\tend"],
)
def test_docx_values_are_xml_safe_without_double_escape(tmp_path, payload):
    source = tmp_path / "template.docx"
    target = tmp_path / "rendered.docx"
    source.write_bytes(_template_bytes())
    template = DocxTemplate(source)
    template.render(
        {"value": payload, "items": ["Một", "Hai"], "enabled": True},
        jinja_env=create_template_environment(),
    )
    template.save(target)

    with zipfile.ZipFile(target) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    text = "".join(root.itertext())
    if "\n" in payload or "\t" in payload:
        assert payload.replace("\n", "").replace("\t", "") in text
        assert b"<w:br" in xml
        assert b"<w:tab" in xml
    else:
        assert payload in text
    assert "MộtHai" in text
    assert "YES" in text
    assert "&amp;amp;" not in xml.decode("utf-8")
    if "</w:t>" in payload:
        assert len(root.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r")) < 10


def test_docx_environment_keeps_strict_undefined(tmp_path):
    source = tmp_path / "template.docx"
    source.write_bytes(_template_bytes())
    template = DocxTemplate(source)
    with pytest.raises(UndefinedError):
        template.render({"items": [], "enabled": False}, jinja_env=create_template_environment())
