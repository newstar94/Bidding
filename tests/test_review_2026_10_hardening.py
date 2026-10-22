"""Regression tests for the 2026-10 review fixes and new deadline/holiday/preview features."""

import asyncio
import io
import json
import sqlite3
import time
from types import SimpleNamespace

import pytest

from backend.auth import auth_routes, auth_service, otp_routes
from backend.auth.otp_security import hash_registration_otp


# --------------------------------------------------------------------------
# Rate-limit peek and login lockout
# --------------------------------------------------------------------------

class _SqliteDatabase:
    def __init__(self, path):
        self.path = path

    def get_connection(self):
        connection = sqlite3.connect(self.path)
        connection.row_factory = sqlite3.Row
        connection.create_function("LEAST", 2, min)
        return _Connection(connection)


class _Connection:
    def __init__(self, connection):
        self.connection = connection
        self.in_transaction = False

    def execute(self, sql, parameters=()):
        return self.cursor().execute(sql, parameters)

    def cursor(self):
        return _Cursor(self.connection.cursor())

    def __getattr__(self, name):
        return getattr(self.connection, name)


class _Cursor:
    def __init__(self, cursor):
        self.cursor = cursor

    def execute(self, sql, parameters=()):
        import re

        sql = re.sub(r"\s+FOR UPDATE(?: OF [\w, ]+)?\s*$", "", sql.strip())
        self.cursor.execute(sql, parameters)
        return self

    def __getattr__(self, name):
        return getattr(self.cursor, name)


@pytest.fixture
def auth_db(tmp_path, monkeypatch):
    database = _SqliteDatabase(tmp_path / "review.sqlite")
    connection = database.get_connection()
    connection.connection.executescript(
        """
        CREATE TABLE rate_limit_buckets (
            bucket_key TEXT PRIMARY KEY, window_started_at INTEGER,
            attempt_count INTEGER, expires_at INTEGER
        );
        CREATE TABLE tai_khoan (
            id TEXT PRIMARY KEY, ten_dang_nhap TEXT, username_norm TEXT, mat_khau TEXT,
            ho_ten TEXT, vai_tro TEXT, email TEXT, email_norm TEXT, da_xac_minh INTEGER,
            ma_xac_minh TEXT, han_xac_minh INTEGER, registration_verified_at INTEGER,
            trang_thai TEXT DEFAULT 'active'
        );
        CREATE TABLE dinh_danh_ngoai (issuer TEXT, subject TEXT, user_id TEXT, email_norm TEXT);
        """
    )
    connection.commit()
    connection.close()
    monkeypatch.setattr(auth_service, "_get_rate_limit_database", lambda: database)
    monkeypatch.setattr(otp_routes, "database", database)
    monkeypatch.setattr(otp_routes, "ensure_personal_word_workspace", lambda *_args: None)
    return database


def test_rate_limit_peek_blocks_once_bucket_is_full(auth_db):
    for _ in range(auth_service.RATE_LIMIT_MAX):
        assert auth_service.get_rate_limit_decision("bucket").allowed
    peek = auth_service.get_rate_limit_decision("bucket", consume_attempt=False)
    assert peek.allowed is False
    assert peek.remaining == 0
    assert peek.retry_after >= 1


def test_login_without_password_failures_does_not_count_against_account(monkeypatch):
    """Junk requests that fail the bot challenge must not lock the account."""
    consumed = []

    async def database_write(function, *args, **_kwargs):
        if function is auth_routes._load_login_user_with_rate_limit:
            return SimpleNamespace(allowed=True, remaining=1), None
        if function is auth_routes._record_failed_account_attempt:
            consumed.append(args)
        return SimpleNamespace(allowed=True, remaining=1)

    async def challenge_fails(*_args, **_kwargs):
        from starlette.responses import JSONResponse

        return JSONResponse({"code": "BOT_CHALLENGE_REQUIRED"}, status_code=403)

    class Request:
        headers = {}
        client = SimpleNamespace(host="198.51.100.7")

        async def json(self):
            return {"username": "victim", "password": "junk"}

    monkeypatch.setattr(auth_routes, "run_database_write", database_write)
    monkeypatch.setattr(auth_routes, "enforce_turnstile", challenge_fails)
    response = asyncio.run(auth_routes.login_api(Request()))
    assert response.status_code == 403
    assert consumed == []


def test_locked_account_requires_challenge_instead_of_hard_block(monkeypatch):
    seen = {}

    async def database_write(function, *_args, **_kwargs):
        if function is auth_routes._load_login_user_with_rate_limit:
            return SimpleNamespace(allowed=False, remaining=0, retry_after=40), None
        return SimpleNamespace(allowed=True, remaining=4)

    async def challenge(_request, _data, *, expected_action, required):
        seen["required"] = required
        from starlette.responses import JSONResponse

        return JSONResponse({"code": "BOT_CHALLENGE_REQUIRED"}, status_code=403)

    class Request:
        headers = {}
        client = SimpleNamespace(host="198.51.100.8")

        async def json(self):
            return {"username": "owner", "password": "pw"}

    monkeypatch.setattr(auth_routes, "run_database_write", database_write)
    monkeypatch.setattr(auth_routes, "enforce_turnstile", challenge)
    monkeypatch.setattr(auth_routes, "_turnstile_escape_available", lambda: True)
    assert asyncio.run(auth_routes.login_api(Request())).status_code == 403
    assert seen["required"] is True

    monkeypatch.setattr(auth_routes, "_turnstile_escape_available", lambda: False)
    assert asyncio.run(auth_routes.login_api(Request())).status_code == 429


# --------------------------------------------------------------------------
# Registration OTP hardening
# --------------------------------------------------------------------------

def _insert_unverified(database, *, user_id="u-1", code="123456", expiry=None, username="alice", email="a@example.com"):
    connection = database.get_connection()
    connection.execute(
        """INSERT INTO tai_khoan (id, ten_dang_nhap, username_norm, mat_khau, ho_ten, vai_tro,
               email, email_norm, da_xac_minh, ma_xac_minh, han_xac_minh)
           VALUES (?, ?, ?, 'hash', 'Alice', 'user', ?, ?, 0, ?, ?)""",
        (user_id, username, username, email, email, hash_registration_otp(code, user_id),
         int(time.time()) + 600 if expiry is None else expiry),
    )
    connection.commit()
    connection.close()


def test_registration_otp_is_invalidated_after_repeated_wrong_codes(auth_db):
    _insert_unverified(auth_db)
    for attempt in range(otp_routes.REGISTRATION_OTP_MAX_FAILURES):
        response = otp_routes._confirm_registration_otp("alice", "000000")
        assert response.status_code == 400
    assert json.loads(response.body)["code"] == "REGISTRATION_OTP_INVALIDATED"
    # Even the correct code no longer works once the code was invalidated.
    assert otp_routes._confirm_registration_otp("alice", "123456").status_code == 400


def test_correct_registration_otp_still_verifies_and_resets_failures(auth_db):
    _insert_unverified(auth_db)
    otp_routes._confirm_registration_otp("alice", "000000")
    assert otp_routes._confirm_registration_otp("alice", "123456") is None
    connection = auth_db.get_connection()
    assert connection.execute("SELECT da_xac_minh FROM tai_khoan").fetchone()[0] == 1
    connection.close()


def test_abandoned_unverified_registration_can_be_reclaimed(auth_db):
    stale = int(time.time()) - otp_routes.UNVERIFIED_ACCOUNT_RELEASE_SECONDS - 10
    _insert_unverified(auth_db, expiry=stale)
    rejection = otp_routes._create_unverified_account(
        username="alice", password_hash="new-hash", name="Owner", role="user",
        email="a@example.com", code="654321",
    )
    assert rejection is None
    connection = auth_db.get_connection()
    rows = connection.execute("SELECT id, mat_khau, ho_ten FROM tai_khoan").fetchall()
    connection.close()
    assert [tuple(row) for row in rows] == [("u-1", "new-hash", "Owner")]


def test_recent_unverified_registration_still_conflicts(auth_db):
    _insert_unverified(auth_db)
    rejection = otp_routes._create_unverified_account(
        username="alice", password_hash="new-hash", name="Other", role="user",
        email="other@example.com", code="654321",
    )
    assert rejection.status_code == 409


# --------------------------------------------------------------------------
# Session bootstrap
# --------------------------------------------------------------------------

def test_index_bootstrap_failure_is_not_reported_as_signed_out(monkeypatch):
    import backend.app as app_module

    async def failing_read(*_args, **_kwargs):
        raise TimeoutError("lock timeout")

    monkeypatch.setattr(app_module, "run_database_read", failing_read)
    monkeypatch.setattr(app_module, "log_error", lambda *_args, **_kwargs: None)
    captured = {}
    original = app_module._workspace_preload_tag

    def capture(session_bootstrap, request_path="/workspace"):
        captured["bootstrap"] = session_bootstrap
        return original(session_bootstrap, request_path)

    monkeypatch.setattr(app_module, "_workspace_preload_tag", capture)
    from starlette.requests import Request

    request = Request({
        "type": "http", "method": "GET", "path": "/tong-quan", "headers": [],
        "query_string": b"", "scheme": "http", "server": ("testserver", 80),
        "client": ("127.0.0.1", 1),
    })
    response = asyncio.run(app_module.index(request))
    assert response.status_code == 200
    assert captured["bootstrap"] == {"valid": None, "reason": "bootstrap_unavailable"}


def test_personal_workspace_bootstrap_skips_lock_when_scope_exists():
    from backend.documents import word_defaults

    statements = []

    class Cursor:
        def execute(self, sql, _params=()):
            statements.append(" ".join(sql.split()))
            self.row = (1,) if "FROM sync_metadata" in sql else None
            return self

        def fetchone(self):
            return self.row

    word_defaults.ensure_personal_word_workspace(Cursor(), "user-a")
    assert not any("pg_advisory_xact_lock" in statement for statement in statements)
    assert not any(statement.startswith("INSERT") for statement in statements)


# --------------------------------------------------------------------------
# Pool sizing, sync tombstones and cache headers
# --------------------------------------------------------------------------

def test_default_pool_covers_database_lanes():
    from backend.db.db_helper import default_pool_max_size

    assert default_pool_max_size({}) == 14
    assert default_pool_max_size({"DATABASE_READ_WORKERS": "4", "DATABASE_WRITE_WORKERS": "2"}) == 8


def test_timestamp_delta_with_too_many_tombstones_requires_full_sync():
    from backend.sync import read_service

    class Scope:
        pass

    class Cursor:
        def execute(self, _sql, params):
            self.limit = params[-1]
            return self

        def fetchall(self):
            return [("goithau", str(index), 1, "2026-01-01") for index in range(self.limit)]

    original = read_service.scoped_deletion_branches
    read_service.scoped_deletion_branches = lambda *_args: [
        ("goithau", "goi_thau", SimpleNamespace(sql="1=1", parameters=()))
    ]
    try:
        with pytest.raises(read_service.TombstoneWindowExceeded):
            read_service._load_visible_deletions(Cursor(), Scope(), since="2026-01-01 00:00:00", limit=3)
    finally:
        read_service.scoped_deletion_branches = original


def test_reference_endpoints_are_cacheable_and_api_stays_no_store():
    from backend.http_middleware import _reference_data_max_age

    assert _reference_data_max_age("/api/address/provinces") == 86_400
    assert _reference_data_max_age("/api/address/wards/01") == 86_400
    assert _reference_data_max_age("/api/holidays") == 300
    assert _reference_data_max_age("/api/get-all-data") == 0


# --------------------------------------------------------------------------
# Holiday calendar store
# --------------------------------------------------------------------------

def test_holiday_store_validates_saves_and_reloads(tmp_path, monkeypatch):
    from backend.holidays import store

    path = tmp_path / "holidays.json"
    path.write_text(json.dumps({"2026": {"holidays": ["2026-01-01"], "working_weekends": []}}), encoding="utf-8")
    monkeypatch.setenv("HOLIDAYS_FILE", str(path))
    assert store.load_holidays()["2026"]["holidays"] == ["2026-01-01"]
    saved = store.save_holidays({"2027": {"holidays": ["2027-02-08", "2027-01-01", "2027-01-01"]}})
    assert saved == {"2027": {"holidays": ["2027-01-01", "2027-02-08"], "working_weekends": []}}
    assert store.load_holidays() == saved
    with pytest.raises(store.HolidayValidationError):
        store.save_holidays({"2027": {"holidays": ["2026-12-31"]}})
    with pytest.raises(store.HolidayValidationError):
        store.save_holidays({"2027": {"working_weekends": ["2027-01-04"]}})  # a Monday
    assert store.calendar_coverage(saved, __import__("datetime").date(2027, 5, 1))["missingYears"] == [2028]


# --------------------------------------------------------------------------
# Deadline center
# --------------------------------------------------------------------------

def test_deadline_items_and_calendar_export():
    from backend.deadlines.routes import build_deadline_calendar, normalize_deadline_item

    item = normalize_deadline_item({
        "targetType": "package", "id": "pkg-1", "maGoiThau": "IB2600000001",
        "tenGoiThau": "Mua sắm, thiết bị; năm 2026", "alertKey": "closingSoon",
        "deadline": "2026-11-03 09:00",
    })
    assert item["dueDate"] == "2026-11-03"
    assert item["route"] == "/goi-thau-chi-tiet/pkg-1"
    calendar = build_deadline_calendar([item], public_url="https://example.test", now=0)
    assert "DTSTART;VALUE=DATE:20261103" in calendar
    assert "SUMMARY:[IB2600000001] Sắp đóng thầu" in calendar
    assert "Mua sắm\\, thiết bị\\; năm 2026" in calendar
    assert all(len(line.encode("utf-8")) <= 75 for line in calendar.split("\r\n"))


# --------------------------------------------------------------------------
# Word preview
# --------------------------------------------------------------------------

def test_docx_preview_escapes_text_and_renders_tables():
    from docx import Document

    from backend.documents.docx_preview import build_document_preview

    document = Document()
    document.add_paragraph("<script>alert(1)</script>")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    buffer = io.BytesIO()
    document.save(buffer)
    preview = build_document_preview(buffer.getvalue(), "x.docx")
    assert preview[0]["filename"] == "x.docx"
    assert "&lt;script&gt;" in preview[0]["html"]
    assert "<script>" not in preview[0]["html"]
    assert "<table" in preview[0]["html"]
