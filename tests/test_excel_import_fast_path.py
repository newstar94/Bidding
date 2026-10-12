import subprocess
import sys
import zipfile
from io import BytesIO

from openpyxl import Workbook
import pytest

from backend.documents.excel_handler import parse_excel
from backend.documents.document_worker import run_document_job


def _workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_xlsx_parser_module_does_not_eagerly_import_pandas():
    completed = subprocess.run(
        [
            sys.executable,
            "-c",
            (
                "import sys; "
                "import backend.documents.excel_handler; "
                "raise SystemExit(1 if 'pandas' in sys.modules else 0)"
            ),
        ],
        check=False,
        capture_output=True,
        text=True,
        timeout=15,
    )
    assert completed.returncode == 0, completed.stderr


def test_xlsx_expert_import_preserves_horizontal_contract():
    content = _workbook_bytes([
        ["Họ tên", "Số CCCD", "Số chứng chỉ"],
        ["Dòng kiểm thử A", "012345678901", "CC-001"],
    ])

    assert parse_excel(content, "chuyengia", kind="xlsx") == [{
        "rowIdx": 2,
        "data": {
            "hoTen": "Dòng kiểm thử A",
            "soCCCD": "012345678901",
            "ngayCapCCCD": "",
            "noiCapCCCD": "",
            "soChungChi": "CC-001",
            "ngayCapChungChi": "",
            "donViCapChungChi": "",
            "anhChungChi": "",
        },
        "isValid": True,
        "comments": "",
    }]


def test_xlsx_expert_import_preserves_vertical_contract():
    content = _workbook_bytes([
        ["Họ tên", "Dòng kiểm thử A", "Dòng kiểm thử B"],
        ["Số CCCD", "012345678901", "123456789012"],
        ["Số chứng chỉ", "CC-001", "CC-002"],
    ])

    rows = parse_excel(content, "chuyengia", kind="xlsx")

    assert [row["rowIdx"] for row in rows] == [1, 2]
    assert [row["data"]["hoTen"] for row in rows] == [
        "Dòng kiểm thử A",
        "Dòng kiểm thử B",
    ]
    assert all(row["isValid"] for row in rows)


def test_isolated_xlsx_worker_preserves_expert_rows(tmp_path):
    content = _workbook_bytes([
        ["Họ tên", "Số CCCD", "Số chứng chỉ"],
        ["Dòng kiểm thử", "123456789012", "CC-WORKER"],
    ])
    workbook_path = tmp_path / "experts.xlsx"
    workbook_path.write_bytes(content)

    rows = run_document_job(
        "parse_excel",
        {
            "content_path": str(workbook_path),
            "kind": "xlsx",
            "import_type": "chuyengia",
        },
        timeout_seconds=10,
    )

    assert len(rows) == 1
    assert rows[0]["isValid"] is True
    assert rows[0]["data"]["hoTen"] == "Dòng kiểm thử"


def _worksheet_bytes(xml):
    original = _workbook_bytes([["Họ tên", "Số CCCD"], ["A", "012345678901"]])
    stream = BytesIO()
    with zipfile.ZipFile(BytesIO(original)) as source, zipfile.ZipFile(stream, "w") as target:
        for name in source.namelist():
            target.writestr(name, xml if name == "xl/worksheets/sheet1.xml" else source.read(name))
    return stream.getvalue()


@pytest.mark.parametrize("rows", [
    '<row r="300"><c r="A300" t="inlineStr"><is><t>A</t></is></c></row>',
    '<row r="1"><c r="ZZZ1" t="inlineStr"><is><t>A</t></is></c></row>',
    '<row r="2"><c r="A2"/></row><row r="2"><c r="B2"/></row>',
    '<row r="2"><c r="A2"/></row><row r="1"><c r="A1"/></row>',
    '<row r="1"><c r="A0"/></row>',
    '<row r="-1"><c r="A1"/></row>',
    '<row r="1048576"><c r="XFD1048576" t="inlineStr"><is><t>A</t></is></c></row>',
    '<row r="999999999999999999999999"><c r="A1"/></row>',
    '<row r="1"><c r="A1suffix"/></row>',
    '<row r="1"><c r="A1"/><c r="a1"/></row>',
    '<row r="1"><c r="A2"/></row>',
])
def test_sparse_or_invalid_coordinates_are_rejected_before_materializing(monkeypatch, rows):
    monkeypatch.setenv("EXCEL_MAX_IMPORT_ROWS", "100")
    content = _worksheet_bytes(
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<sheetData>{rows}</sheetData></worksheet>'
    )
    with pytest.raises(ValueError):
        parse_excel(content, "chuyengia", kind="xlsx")


def test_rectangular_budget_counts_sparse_width_times_height(monkeypatch):
    monkeypatch.setenv("EXCEL_MAX_IMPORT_ROWS", "100")
    content = _worksheet_bytes(
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData><row r="100"><c r="CV100" t="inlineStr"><is><t>A</t></is></c></row>'
        '</sheetData></worksheet>'
    )
    from backend.documents.xlsx_fast_reader import read_first_worksheet_rows

    with pytest.raises(ValueError, match="ô"):
        read_first_worksheet_rows(content, max_rows=101, max_columns=101, max_cells=1000)


@pytest.mark.parametrize("vertical", [False, True])
def test_configured_record_boundary_preserves_both_orientations(monkeypatch, vertical):
    monkeypatch.setenv("EXCEL_MAX_IMPORT_ROWS", "100")
    names = [f"Người {index}" for index in range(100)]
    numbers = [f"{index:012d}" for index in range(100)]
    rows = [["Họ tên", *names], ["Số CCCD", *numbers]] if vertical else [
        ["Họ tên", "Số CCCD"], *[[name, number] for name, number in zip(names, numbers)],
    ]
    assert len(parse_excel(_workbook_bytes(rows), "chuyengia", kind="xlsx")) == 100


@pytest.mark.parametrize("formatted_row", [257, 1_048_576])
def test_empty_trailing_formatting_preserves_valid_import(monkeypatch, formatted_row):
    from openpyxl.styles import Font

    monkeypatch.setenv("EXCEL_MAX_IMPORT_ROWS", "100")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Họ tên", "Số CCCD"])
    sheet.append(["Một người", "012345678901"])
    sheet.cell(formatted_row, 1).font = Font(bold=True)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    result = parse_excel(stream.getvalue(), "chuyengia", kind="xlsx")
    assert len(result) == 1
    assert result[0]["data"]["hoTen"] == "Một người"
