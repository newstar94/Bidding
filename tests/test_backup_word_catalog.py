import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docxtpl import DocxTemplate

from backend.documents.template_catalog.storage import ImmutableTemplateStorage
from backend.observability.backup_validation import sha256_file
from scripts import backup


def _snapshot(monkeypatch, tmp_path):
    catalog = tmp_path / "catalog-source"
    storage = ImmutableTemplateStorage(catalog)
    document = Document()
    document.add_paragraph("Đơn vị: {{ company_name }}")
    template_bytes = BytesIO()
    document.save(template_bytes)
    key, digest, _size = storage.put("org-1", template_bytes.getvalue())
    for name in ("uploads-source", "templates-source"):
        (tmp_path / name).mkdir()
    monkeypatch.setenv("WORD_TEMPLATE_CATALOG_ENABLED", "true")
    monkeypatch.setenv("BIDDING_WORD_TEMPLATE_CATALOG_DIR", str(catalog))
    monkeypatch.setenv("BACKUP_DATABASE_URL", "postgresql://backup:test@localhost/test")
    monkeypatch.setattr(backup, "_postgres_binary", lambda name: name)

    def dump(command, **_kwargs):
        Path(command[command.index("--file") + 1]).write_bytes(b"fake database dump")
        return SimpleNamespace(returncode=0, stderr="")

    monkeypatch.setattr(backup.subprocess, "run", dump)
    assert backup.cmd_create(SimpleNamespace(
        backup_dir=str(tmp_path / "snapshots"), uploads=str(tmp_path / "uploads-source"),
        word_templates=str(tmp_path / "templates-source"),
    )) == 0
    snapshot = next((tmp_path / "snapshots").iterdir())
    return snapshot, key, digest


@pytest.mark.parametrize("restore_succeeds", [True, False])
def test_backup_catalog_bytes_and_restore_preserve_immutable_digest(monkeypatch, tmp_path, restore_succeeds):
    snapshot, key, digest = _snapshot(monkeypatch, tmp_path)
    manifest = json.loads((snapshot / "manifest.json").read_text(encoding="utf-8"))
    entry = next((entry for entry in manifest["files"] if entry["relativePath"] == f"word-catalog/{key}"), None)
    assert entry is not None, "backup omitted the catalog-only template"
    assert entry["sha256"] == digest == sha256_file(snapshot / "word-catalog" / key)
    assert backup.cmd_verify(SimpleNamespace(snapshot=str(snapshot), require_complete=True)) == 0

    destinations = {name: tmp_path / f"restore-{name}" for name in ("uploads", "word-templates", "word-catalog")}
    for name, destination in destinations.items():
        destination.mkdir()
        (destination / "newer.txt").write_bytes(b"preserve if database restore fails")
        monkeypatch.setenv({"uploads": "BIDDING_UPLOAD_DIR", "word-templates": "BIDDING_WORD_TEMPLATE_DIR",
                            "word-catalog": "BIDDING_WORD_TEMPLATE_CATALOG_DIR"}[name], str(destination))
    monkeypatch.setenv("DATABASE_URL", "postgresql://restore:test@localhost/isolated_restore")

    def restore(command, **_kwargs):
        assert "--single-transaction" in command
        assert "isolated_restore" in command
        assert ImmutableTemplateStorage(destinations["word-catalog"]).read("org-1", key, digest) == (snapshot / "word-catalog" / key).read_bytes()
        return SimpleNamespace(returncode=0 if restore_succeeds else 1, stderr="mock failure")

    monkeypatch.setattr(backup.subprocess, "run", restore)
    assert backup.cmd_restore(SimpleNamespace(snapshot=str(snapshot))) == (0 if restore_succeeds else 1)
    if restore_succeeds:
        restored = ImmutableTemplateStorage(destinations["word-catalog"]).read("org-1", key, digest)
        template = DocxTemplate(BytesIO(restored))
        template.render({"company_name": "Công ty khôi phục"})
        output = BytesIO()
        template.save(output)
        assert Document(BytesIO(output.getvalue())).paragraphs[0].text == "Đơn vị: Công ty khôi phục"
    else:
        for destination in destinations.values():
            assert (destination / "newer.txt").read_bytes() == b"preserve if database restore fails"


def test_catalog_backup_refuses_changed_bytes_during_dump(monkeypatch, tmp_path):
    catalog = tmp_path / "catalog"
    key, _digest, _size = ImmutableTemplateStorage(catalog).put("org-1", b"first")
    for name in ("uploads", "templates"):
        (tmp_path / name).mkdir()
    monkeypatch.setenv("WORD_TEMPLATE_CATALOG_ENABLED", "true")
    monkeypatch.setenv("BIDDING_WORD_TEMPLATE_CATALOG_DIR", str(catalog))
    monkeypatch.setenv("BACKUP_DATABASE_URL", "postgresql://backup:test@localhost/test")
    monkeypatch.setattr(backup, "_postgres_binary", lambda name: name)

    def dump(command, **_kwargs):
        Path(command[command.index("--file") + 1]).write_bytes(b"fake dump")
        (catalog / key).write_bytes(b"changed")
        return SimpleNamespace(returncode=0, stderr="")

    monkeypatch.setattr(backup.subprocess, "run", dump)
    assert backup.cmd_create(SimpleNamespace(
        backup_dir=str(tmp_path / "snapshots"), uploads=str(tmp_path / "uploads"),
        word_templates=str(tmp_path / "templates"),
    )) == 1
    assert list((tmp_path / "snapshots").iterdir()) == []
